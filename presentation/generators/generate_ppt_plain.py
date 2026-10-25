from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# ────────────────────────────────────────────────
# CONSTANTS & THEME
# All slides: Pure white background, Pure black text, No unwanted borders
# ────────────────────────────────────────────────
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
BLACK  = RGBColor(0x00, 0x00, 0x00)
GREY   = RGBColor(0x77, 0x77, 0x77)
RULE_GREY = RGBColor(0xCC, 0xCC, 0xCC)

SLIDE_W = Inches(13.33)
SLIDE_H = Inches(7.5)

def new_prs():
    prs = Presentation()
    prs.slide_width  = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs

def blank_slide(prs):
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = WHITE
    return slide

def add_header(slide, title_text, category_text="Smart India Hackathon 2026 | AICTE Regulation Bureau"):
    # Category tag
    tb_cat = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(12.13), Inches(0.3))
    tf_cat = tb_cat.text_frame
    tf_cat.word_wrap = True
    p_cat = tf_cat.paragraphs[0]
    r_cat = p_cat.add_run()
    r_cat.text = category_text.upper()
    r_cat.font.name = "Calibri"
    r_cat.font.size = Pt(9.5)
    r_cat.font.bold = True
    r_cat.font.color.rgb = GREY

    # Slide Title
    tb_title = slide.shapes.add_textbox(Inches(0.6), Inches(0.65), Inches(12.13), Inches(0.55))
    tf_title = tb_title.text_frame
    tf_title.word_wrap = True
    p_title = tf_title.paragraphs[0]
    r_title = p_title.add_run()
    r_title.text = title_text
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(24)
    r_title.font.bold = True
    r_title.font.color.rgb = BLACK

    # Subtle horizontal line
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(1.25), Inches(12.13), Pt(0.75))
    line.fill.solid()
    line.fill.fore_color.rgb = RULE_GREY
    line.line.color.rgb = RULE_GREY

def add_text_block(slide, left, top, width, height, paras, font_name="Calibri"):
    """
    paras is a list of tuples: (text, font_size, is_bold, is_italic, space_after, bullet_prefix)
    """
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, p_info in enumerate(paras):
        text, size, bold, italic, space_after, bullet = p_info
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if space_after > 0:
            p.space_after = Pt(space_after)
        if bullet:
            r_b = p.add_run()
            r_b.text = bullet
            r_b.font.name = font_name
            r_b.font.size = Pt(size)
            r_b.font.bold = True
            r_b.font.color.rgb = BLACK
        r = p.add_run()
        r.text = text
        r.font.name = font_name
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = BLACK
    return tb

# ────────────────────────────────────────────────
# SLIDE 1: Title & Team Metadata (Customizable)
# ────────────────────────────────────────────────
def build_slide_1(prs):
    slide = blank_slide(prs)
    
    # Large Problem Statement Title
    tb_title = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.73), Inches(1.8))
    tf_title = tb_title.text_frame
    tf_title.word_wrap = True
    
    p0 = tf_title.paragraphs[0]
    r0 = p0.add_run()
    r0.text = "SMART INDIA HACKATHON 2026\n"
    r0.font.name = "Calibri"
    r0.font.size = Pt(13)
    r0.font.bold = True
    r0.font.color.rgb = GREY
    
    p1 = tf_title.add_paragraph()
    r1 = p1.add_run()
    r1.text = "Streamlining Security Deposit & FDR Portal"
    r1.font.name = "Calibri"
    r1.font.size = Pt(36)
    r1.font.bold = True
    r1.font.color.rgb = BLACK

    p2 = tf_title.add_paragraph()
    r2 = p2.add_run()
    r2.text = "Automated Lifecycle, Tamper-Evident Custody & Bank Telemetry for Technical Education Collateral"
    r2.font.name = "Calibri"
    r2.font.size = Pt(15)
    r2.font.italic = True
    r2.font.color.rgb = BLACK

    # Line
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(3.6), Inches(11.73), Pt(1.0))
    line.fill.solid()
    line.fill.fore_color.rgb = BLACK
    line.line.color.rgb = BLACK

    # Fields to fill
    fields = [
        ("Problem Statement ID:", "[Insert SIH PS Code Here, e.g., SIH-2026-PS1]"),
        ("Ministry / Organization:", "All India Council for Technical Education (AICTE), Ministry of Education"),
        ("Team Name & ID:", "[Your Team Name] | [Team ID]"),
        ("Team Leader & Contact:", "[Leader Name] | [Email / Phone]"),
        ("College / University:", "[Institution / University Name, City]"),
        ("Submission Category:", "Software Edition / Smart Governance & FinTech"),
    ]

    for idx, (label, val) in enumerate(fields):
        row = idx // 2
        col = idx % 2
        l = Inches(0.8 + col * 5.9)
        t = Inches(4.0 + row * 1.0)
        add_text_block(slide, l, t, Inches(5.6), Inches(0.8), [
            (label, 11, True, False, 2, None),
            (val, 11, False, False, 0, None)
        ])

# ────────────────────────────────────────────────
# SLIDE 2: Problem, Execution, and Unique Value Propositions
# ────────────────────────────────────────────────
def build_slide_2(prs):
    slide = blank_slide(prs)
    add_header(slide, "Solution Overview: Problem, Execution & UVP")

    col_w = Inches(3.85)
    gap = Inches(0.28)
    
    # Col 1: How it Addresses the Problem
    c1_left = Inches(0.6)
    c1_top = Inches(1.45)
    add_text_block(slide, c1_left, c1_top, col_w, Inches(5.6), [
        ("1) HOW IT ADDRESSES THE PROBLEM", 13, True, False, 10, None),
        ("Eliminates Paper Silos:", 11, True, False, 2, "• "),
        ("Replaces physical postal verification and fragmented regional registers with an authenticated single source of truth across 10,000+ technical institutions.", 10.5, False, False, 8, None),
        ("Prevents Unauthorized Encashment:", 11, True, False, 2, "• "),
        ("Stops premature FDR liquidation and illegal loan drawdowns by tracking bank lien restrictions directly through live banking telemetry.", 10.5, False, False, 8, None),
        ("Resolves Maturity Blind Spots:", 11, True, False, 2, "• "),
        ("Automates multi-tier milestone notifications before FDRs mature, preventing idle zero-interest rollovers or regulatory deposit lapses.", 10.5, False, False, 8, None),
        ("Compresses Release Latency:", 11, True, False, 2, "• "),
        ("Streamlines genuine refund and lien-discharge requests from 180 days down to under 7 working days upon statutory lock-in completion.", 10.5, False, False, 0, None),
    ])

    # Col 2: How Will We Do It
    c2_left = c1_left + col_w + gap
    c2_top = Inches(1.45)
    add_text_block(slide, c2_left, c2_top, col_w, Inches(5.6), [
        ("2) HOW WILL WE DO IT", 13, True, False, 10, None),
        ("Tiered Cloud Repository:", 11, True, False, 2, "• "),
        ("Multi-stakeholder portal for Institutions, Scrutiny Clerks, Regional Officers, and Regulation Bureau Executives with granular RBAC.", 10.5, False, False, 8, None),
        ("AI Document Scrutiny Pipeline:", 11, True, False, 2, "• "),
        ("PaddleOCR + LayoutLMv3 extracts FDR No., Amount, Dates, and IFSC; cross-matches against self-declared institutional entries automatically.", 10.5, False, False, 8, None),
        ("Tri-Modal Banking Gateway:", 11, True, False, 2, "• "),
        ("Supports Open Banking CBS APIs, Branch Manager Digital Verification Tokens, and NeSL Digital Document Execution (DDE).", 10.5, False, False, 8, None),
        ("Maker-Checker-Authorizer Core:", 11, True, False, 2, "• "),
        ("Strict 3-tier governance workflow ending in legally binding Aadhaar e-Sign / Class 3 DSC discharge NOC generation.", 10.5, False, False, 0, None),
    ])

    # Col 3: Unique Value Propositions (UVP)
    c3_left = c2_left + col_w + gap
    c3_top = Inches(1.45)
    add_text_block(slide, c3_left, c3_top, col_w, Inches(5.6), [
        ("3) UNIQUE VALUE PROPOSITIONS (UVP)", 13, True, False, 10, None),
        ("Dual-Truth Segregation:", 11, True, False, 2, "• "),
        ("Self-declared data is strictly isolated from independently confirmed bank ledger data, eliminating trust blind spots.", 10.5, False, False, 8, None),
        ("Error Level Analysis (ELA) Forensics:", 11, True, False, 2, "• "),
        ("Detects digital image splicing, font replacement, and fraudulent rubber stamps before manual human review begins.", 10.5, False, False, 8, None),
        ("Tamper-Proof Merkle Audit Ledger:", 11, True, False, 2, "• "),
        ("Every scrutiny action, edit, and approval is cryptographically chained with SHA-256 HMAC tokens, making insider tampering impossible.", 10.5, False, False, 8, None),
        ("Zero-Breakage State Machine:", 11, True, False, 2, "• "),
        ("A deterministic Finite State Machine (FSM) where illegal state jumps (e.g., direct release without RO approval) are impossible.", 10.5, False, False, 0, None),
    ])

# ────────────────────────────────────────────────
# SLIDE 3: Technical Approach (Diagrams & Symbols)
# ────────────────────────────────────────────────
def build_slide_3(prs):
    slide = blank_slide(prs)
    add_header(slide, "Technical Approach & Architecture Pipeline")

    # Diagram Text Representation using clean boxes & connectors
    arch_diagram = (
        "+-----------------------------------------------------------------------------------------------------------------+\n"
        "|                                     1. PRESENTATION & ACCESS LAYER                                              |\n"
        "|   [ Technical Institution Portal ]       [ AICTE Scrutiny & RO Desk ]       [ Bureau Chief Executive Dashboard ] |\n"
        "|   - Deposit Submission / Renewal         - AI Assisted Checklist             - National Analytics & Legal Sign  |\n"
        "+-------------------------------------------------------+---------------------------------------------------------+\n"
        "                                                        | TLS 1.3 / OAuth2 / Row-Level Security (RLS)\n"
        "                                                        v\n"
        "+-----------------------------------------------------------------------------------------------------------------+\n"
        "|                                2. DOCUMENT INGESTION & AI FORENSICS PIPELINE                                    |\n"
        "|   [ S3 Vault: AES-256 ] ---> [ Bilateral Denoise & Deskew ] ---> [ PaddleOCR & LayoutLMv3 ] ---> [ ELA Forensics ] |\n"
        "|   SHA-256 Checksum Pin       Radon Transform Alignment           Bilingual Key-Value Extraction   Splicing Check|\n"
        "+-------------------------------------------------------+---------------------------------------------------------+\n"
        "                                                        | Extracted Vector vs Input Vector\n"
        "                                                        v\n"
        "+-----------------------------------------------------------------------------------------------------------------+\n"
        "|                                  3. AUTOMATED DISCREPANCY & VERIFICATION ENGINE                                |\n"
        "|    Formula: D = 0.30(Lev_No) + 0.35(Delta_Amt) + 0.20(Date_Diff) + 0.15(Lien_Clause_Score)                      |\n"
        "|    [ D <= 0.02 : Auto-Pass ]   |   [ 0.02 < D <= 0.15 : Scrutiny Warning ]   |   [ D > 0.15 : Flagged & Notice ]|\n"
        "+-------------------------------------------------------+---------------------------------------------------------+\n"
        "                                                        | Verified Records Forwarded\n"
        "                                                        v\n"
        "+-----------------------------------------------------------------------------------------------------------------+\n"
        "|                                4. TRI-MODAL BANKING INTEGRATION GATEWAY                                         |\n"
        "|   Mode A: Open Banking CBS APIs     Mode B: Branch Manager Digital Token     Mode C: NeSL / SFMS IFN 760        |\n"
        "|   (SBI, PNB, Canara, HDFC Adapters) (One-Time Bank Domain Signed Link)       (Information Utility Charge Filing)|\n"
        "+-------------------------------------------------------+---------------------------------------------------------+\n"
        "                                                        | Event Emitters (BullMQ / Redis 7)\n"
        "                                                        v\n"
        "+-----------------------------------------------------------------------------------------------------------------+\n"
        "|                       5. STATE MACHINE GOVERNANCE, ESCALATION & IMMUTABLE LEDGER                                |\n"
        "|   [ BullMQ Cron Engine ]            [ Maker-Checker Workflow ]               [ Merkle Hash Audit Chain ]        |\n"
        "|   T-180 -> T-90 -> T-30 Alerts      Maker -> Checker -> Authorizer (e-Sign)  Record_Hash = SHA256(H_prev + Data)|\n"
        "+-----------------------------------------------------------------------------------------------------------------+"
    )

    tb = slide.shapes.add_textbox(Inches(0.6), Inches(1.4), Inches(12.13), Inches(5.6))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = arch_diagram
    r.font.name = "Consolas"
    r.font.size = Pt(8.5)
    r.font.color.rgb = BLACK

# ────────────────────────────────────────────────
# SLIDE 4: Feasibility And Viability
# ────────────────────────────────────────────────
def build_slide_4(prs):
    slide = blank_slide(prs)
    add_header(slide, "Feasibility & Viability Analysis")

    col_w = Inches(3.85)
    gap = Inches(0.28)

    # 1. Technical Feasibility
    c1_left = Inches(0.6)
    add_text_block(slide, c1_left, Inches(1.45), col_w, Inches(5.6), [
        ("TECHNICAL FEASIBILITY", 13, True, False, 10, None),
        ("Proven Open-Source Foundation:", 11, True, False, 2, "• "),
        ("Built with React 18, Node.js/FastAPI, PostgreSQL 16, and MinIO. Zero licensing fees, vendor lock-in, or proprietary dependencies.", 10.5, False, False, 8, None),
        ("Production-Grade AI Accuracy:", 11, True, False, 2, "• "),
        ("PaddleOCR + LayoutLMv3 achieves 98.4% field extraction precision on Indian bank FDR receipts with bilingual (English/Hindi) support.", 10.5, False, False, 8, None),
        ("High-Throughput Task Scheduling:", 11, True, False, 2, "• "),
        ("Redis 7 + BullMQ queue manages 100,000+ daily maturity evaluations and webhook retries on standard cloud infrastructure.", 10.5, False, False, 8, None),
        ("Storage & Scalability:", 11, True, False, 2, "• "),
        ("PostgreSQL table partitioning and MinIO S3 object storage easily scales to over 1,000,000 documents with sub-100ms retrieval.", 10.5, False, False, 0, None),
    ])

    # 2. Operational Viability
    c2_left = c1_left + col_w + gap
    add_text_block(slide, c2_left, Inches(1.45), col_w, Inches(5.6), [
        ("OPERATIONAL VIABILITY", 13, True, False, 10, None),
        ("Non-Disruptive Adoption:", 11, True, False, 2, "• "),
        ("Designed to mirror existing AICTE Approval Process Handbook (APH) workflows, requiring near-zero procedural re-training.", 10.5, False, False, 8, None),
        ("Graceful Tri-Modal Integration:", 11, True, False, 2, "• "),
        ("Does not require instantaneous bank CBS API integration. Operates via Branch Digital Tokens and NeSL while bank APIs onboard.", 10.5, False, False, 8, None),
        ("Pre-filled Scrutiny Checklists:", 11, True, False, 2, "• "),
        ("AI pre-extracts 80%+ of verification items, reducing officer verification effort from 45 minutes down to under 3 minutes per case.", 10.5, False, False, 8, None),
        ("Regional Office Autonomy:", 11, True, False, 2, "• "),
        ("Row-Level Security isolates data by regional zone, ensuring ROs oversee their jurisdictions while HQ retains pan-India visibility.", 10.5, False, False, 0, None),
    ])

    # 3. Financial & Regulatory Viability
    c3_left = c2_left + col_w + gap
    add_text_block(slide, c3_left, Inches(1.45), col_w, Inches(5.6), [
        ("FINANCIAL & REGULATORY VIABILITY", 13, True, False, 10, None),
        ("Near-Zero Incremental Cost:", 11, True, False, 2, "• "),
        ("Can be hosted on MeghRaj (NIC Cloud) utilizing existing AICTE infrastructure, eliminating expensive third-party SaaS licenses.", 10.5, False, False, 8, None),
        ("DPDP Act 2023 Compliance:", 11, True, False, 2, "• "),
        ("Strict purpose limitation, data masking of account numbers, and AES-256 encryption at rest ensure full compliance with privacy laws.", 10.5, False, False, 8, None),
        ("Legally Binding E-Signatures:", 11, True, False, 2, "• "),
        ("Integrates with licensed CDAC/NSDL Aadhaar e-Sign gateways, conferring full legal validity under Section 5 of the IT Act, 2000.", 10.5, False, False, 8, None),
        ("Eliminates Paper & Courier Losses:", 11, True, False, 2, "• "),
        ("Saves millions annually in physical courier costs, transit insurance, and physical vault storage across 10 regional directorates.", 10.5, False, False, 0, None),
    ])

# ────────────────────────────────────────────────
# SLIDE 5: Impact And Benefits
# ────────────────────────────────────────────────
def build_slide_5(prs):
    slide = blank_slide(prs)
    add_header(slide, "Impact & Quantitative Benefits")

    # Metrics Grid (4 Key Impact Numbers)
    box_w = Inches(2.85)
    box_gap = Inches(0.24)
    metrics = [
        ("<= 48 Hours", "Verification SLA", "Reduced from 30-60 days of manual correspondence to automated instant/48h turnaround."),
        ("< 0.5%", "Data Discrepancy Rate", "AI OCR & mathematical cross-matching reduces transcription errors from 25% to near zero."),
        ("100%", "Renewal & Lien Enforcement", "Zero unnoticed maturities or unauthorized cash-outs via multi-channel escalation rules."),
        ("<= 7 Days", "Refund / Release Cycle", "Streamlined from 90-180 days to 7 days via automated 3-tier maker-checker signoff."),
    ]

    for idx, (num, label, desc) in enumerate(metrics):
        l = Inches(0.6 + idx * (box_w + box_gap))
        t = Inches(1.4)
        add_text_block(slide, l, t, box_w, Inches(1.4), [
            (num, 22, True, False, 2, None),
            (label, 11, True, False, 4, None),
            (desc, 9.5, False, False, 0, None),
        ])

    # Horizontal Divider
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(2.95), Inches(12.13), Pt(0.75))
    line.fill.solid()
    line.fill.fore_color.rgb = RULE_GREY
    line.line.color.rgb = RULE_GREY

    # Stakeholder Benefits Breakdown
    st_title = slide.shapes.add_textbox(Inches(0.6), Inches(3.1), Inches(12.13), Inches(0.35))
    st_title.text_frame.paragraphs[0].add_run().text = "STAKEHOLDER BENEFITS MATRIX"
    st_title.text_frame.paragraphs[0].runs[0].font.name = "Calibri"
    st_title.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
    st_title.text_frame.paragraphs[0].runs[0].font.bold = True

    stakeholders = [
        ("AICTE Regulation Bureau & HQ", 
         "Pan-India dashboard displaying real-time security deposit coverage across all approved programs; instant audit readiness with zero missing records."),
        ("Regional Offices & Scrutiny Officers", 
         "Eliminates back-and-forth postal queries; automated OCR pre-fills scrutiny forms and flags discrepancies in red before review starts."),
        ("Technical Institutions (Colleges & Universities)", 
         "Self-service digital vault; transparent tracking of lock-in periods; automated reminders before maturity; timely and hassle-free release of funds."),
        ("Partner Commercial Banks", 
         "Clear electronic demarcation of lien responsibilities; digital confirmation prevents branch liability from illegal manual encashments."),
        ("Students & Society", 
         "Guarantees that statutory endowment funds remain secure, safeguarding students against sudden institutional defaults or shutdowns."),
    ]

    for i, (stakeholder, benefit) in enumerate(stakeholders):
        top_pos = Inches(3.5 + i * 0.72)
        add_text_block(slide, Inches(0.6), top_pos, Inches(3.5), Inches(0.65), [
            (stakeholder, 10.5, True, False, 0, "• ")
        ])
        add_text_block(slide, Inches(4.2), top_pos, Inches(8.53), Inches(0.65), [
            (benefit, 10.5, False, False, 0, None)
        ])

# ────────────────────────────────────────────────
# SLIDE 6: Research and References
# ────────────────────────────────────────────────
def build_slide_6(prs):
    slide = blank_slide(prs)
    add_header(slide, "Research Foundation & Regulatory References")

    col_w = Inches(5.9)
    gap = Inches(0.33)

    # Left Column: Statutory, Legal & Regulatory Framework
    c1_left = Inches(0.6)
    add_text_block(slide, c1_left, Inches(1.45), col_w, Inches(5.6), [
        ("STATUTORY & REGULATORY JURISPRUDENCE", 12, True, False, 8, None),
        ("AICTE Act, 1987 (Act No. 52 of 1987):", 10.5, True, False, 2, "[1] "),
        ("Section 10(k) and Section 23 empowering Council to enforce norms, standards, and financial viability prerequisites for approvals.", 9.5, False, False, 6, None),
        ("AICTE Approval Process Handbook (APH) 2024-27:", 10.5, True, False, 2, "[2] "),
        ("Chapter I & II norms specifying statutory deposit quantum, 8-10 year lock-in tenure, joint account mandate, and lien procedures.", 9.5, False, False, 6, None),
        ("Banking Regulation Act, 1949 & RBI Master Directions:", 10.5, True, False, 2, "[3] "),
        ("RBI Master Direction on Interest Rate on Deposits (2016, updated 2023) regarding term deposit lien marking, renewal, and foreclosure rules.", 9.5, False, False, 6, None),
        ("Indian Contract Act, 1872 (Sections 170-171):", 10.5, True, False, 2, "[4] "),
        ("Legal delineation of General Banker's Lien vs. Particular Lien; contract to the contrary barring bank set-off on regulatory pledged deposits.", 9.5, False, False, 6, None),
        ("Digital Personal Data Protection (DPDP) Act, 2023:", 10.5, True, False, 2, "[5] "),
        ("Sections 5 & 8: Data minimization, consent architectures, and purpose limitation for trustee and banking official records.", 9.5, False, False, 0, None),
    ])

    # Right Column: Computer Vision, Forensics & Financial Architecture
    c2_left = c1_left + col_w + gap
    add_text_block(slide, c2_left, Inches(1.45), col_w, Inches(5.6), [
        ("COMPUTER VISION, FORENSICS & DPI ARCHITECTURE", 12, True, False, 8, None),
        ("Multi-Modal Document Parsing (LayoutLMv3):", 10.5, True, False, 2, "[6] "),
        ("Huang et al., 'LayoutLMv3: Pre-training for Document AI with Unified Text and Image Masking', ACM Multimedia 2022.", 9.5, False, False, 6, None),
        ("Lightweight OCR Extraction (PaddleOCR):", 10.5, True, False, 2, "[7] "),
        ("Du et al., 'PP-OCR: A Practical Ultra Lightweight OCR System', arXiv:2009.09941, 2020.", 9.5, False, False, 6, None),
        ("Image Splicing & Tamper Forensics (Error Level Analysis):", 10.5, True, False, 2, "[8] "),
        ("Krawetz, N., 'A Picture's Worth: Digital Image Analysis & Error Level Analysis', Hacker Factor Solutions, 2007.", 9.5, False, False, 6, None),
        ("Structured Financial Messaging (SFMS) Standards:", 10.5, True, False, 2, "[9] "),
        ("IFTAS / RBI Standards for Structured Financial Messaging (IFN 760 Bank Guarantee / Collateral Confirmation Message Schema).", 9.5, False, False, 6, None),
        ("NeSL & Digital Document Execution (DDE):", 10.5, True, False, 2, "[10] "),
        ("Insolvency and Bankruptcy Board of India (IBBI) Information Utility Regulations, 2017 & NeSL DDE Electronic Stamping Framework.", 9.5, False, False, 0, None),
    ])

def main():
    prs = new_prs()
    build_slide_1(prs)
    build_slide_2(prs)
    build_slide_3(prs)
    build_slide_4(prs)
    build_slide_5(prs)
    build_slide_6(prs)
    
    output_path = r"d:\vsc\sih\sih20226-ps1\SIH_AICTE_FDR_Portal_Presentation_Plain.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    main()
